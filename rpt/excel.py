"""Build the downloadable Excel workbook from parsed rows."""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
WRAP = Alignment(wrap_text=True, vertical="top")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")       # pale yellow row
FLAG_FONT = Font(color="9C2500")

_NUMERIC_HINTS = ("value", "basic", "sef", "penalty", "total", "share", "tax",
                  "amount", "due", "discount", "fund", "fee", "dsc", "interest",
                  "area", "principal", "net")


def _is_numeric_col(name: str) -> bool:
    low = name.lower()
    if low in ("year", "tax year", "page no", "page no.", "page number",
               "installment no", "qtr", "line"):
        return False
    return any(h in low for h in _NUMERIC_HINTS)


def build_workbook(location: str, columns: list[str], rows: list[dict],
                   notes: list[str], source_files: list[str],
                   flagged: int = 0) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Data"
    ws.append(columns)
    for c in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    ws.freeze_panes = "A2"

    numeric_cols = {i for i, name in enumerate(columns) if _is_numeric_col(name)}
    review_idx = columns.index("Review") if "Review" in columns else -1
    for ridx, r in enumerate(rows, start=2):
        values = []
        for i, col in enumerate(columns):
            v = r.get(col, "")
            if i in numeric_cols and isinstance(v, (int, float)):
                values.append(float(v))
            else:
                values.append("" if v is None else v)
        ws.append(values)
        if review_idx >= 0 and str(r.get("Review", "")).strip():
            for i in range(len(columns)):
                ws.cell(row=ridx, column=i + 1).fill = FLAG_FILL
            ws.cell(row=ridx, column=review_idx + 1).font = FLAG_FONT

    for i in range(len(columns)):
        letter = get_column_letter(i + 1)
        width = max([len(str(columns[i]))] +
                    [len(str(r.get(columns[i], ""))) for r in rows[:200]] + [8])
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 40)
        if i in numeric_cols:
            for row in range(2, len(rows) + 2):
                ws.cell(row=row, column=i + 1).number_format = "#,##0.00"

    info = wb.create_sheet("Summary", 0)
    info["A1"] = f"RPT Assessment extract — {location}"
    info["A1"].font = TITLE_FONT
    meta = [
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Location", location),
        ("Rows extracted", len(rows)),
        ("Rows needing review", f"{flagged}  (highlighted; see the 'Review' column)"),
        ("Source file(s)", ", ".join(source_files) or "—"),
    ]
    row = 3
    for k, v in meta:
        info.cell(row=row, column=1, value=k).font = Font(bold=True)
        info.cell(row=row, column=2, value=v)
        row += 1

    row += 1
    info.cell(row=row, column=1, value="Notes").font = Font(bold=True)
    row += 1
    for n in (notes or ["No warnings."]):
        info.cell(row=row, column=1, value=f"• {n}")
        row += 1
    row += 1
    warn = ("Values are OCR-read from scanned PDFs. Always check figures against "
            "the source document before using them.")
    c = info.cell(row=row, column=1, value=warn)
    c.font = Font(italic=True, color="9C2500")
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 70

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
